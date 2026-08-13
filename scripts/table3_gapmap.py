"""
Manuscript Table 3 — USPSTF evidentiary indicators by substance condition.

Unlike the per-condition heat map in descriptive_analyses.py (which strata by
era), this table puts the three conditions side by side against a pooled "All"
column, which is the form the manuscript prints.

Inputs:
    data/<Condition>_Dataset_Cleaned.xlsx   (output of clean_dataset.py)

Output:
    tables/Table3_USPSTF_GapMap_bySubstance.docx

Usage:
    python scripts/table3_gapmap.py

Every cell is computed from the cleaned frames; nothing is hardcoded. Missing
or not-reported counts as not met, per the strict USPSTF gap convention used
throughout the analysis.
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt

_BASE    = Path(__file__).parent.parent
_SCRIPTS = Path(__file__).parent
sys.path.insert(0, str(_SCRIPTS))

import descriptive_analyses as da  # noqa: E402  (docx styling + % helpers)

CONDITIONS = [('Alcohol', 'Alcohol'), ('Opioid', 'Opioids'), ('Smoking', 'Smoking')]

OUTPUT = _BASE / 'tables' / 'Table3_USPSTF_GapMap_bySubstance.docx'

# (domain, label, column, kind). Kinds beyond descriptive_analyses' three:
#   'mean'          — mean of a percentage column across studies reporting it
#   'us'            — % of studies whose country is the US
#   'median'        — median of a numeric column, printed without a % sign
#   'representative'— the composite recruitment filter defined in the manuscript
#                     footnote: ≤65% male, ≥40% BIPOC, low-SES participants included
INDICATORS = [
    ('Study Identification & Design', 'Were subjects randomized?',                        'RandomizationYN',            'binary'),
    ('Study Identification & Design', 'Were participants blinded?',                       'BlindingYN',                 'binary'),
    ('Study Identification & Design', 'Was the study double-blinded?',                    'Double+Blind',               'binary'),
    ('Study Identification & Design', 'Did the study include a control or comparison group?*', 'Comparator',            'binary'),
    ('Study Identification & Design', 'Interfaces with primary care',                     'PrimaryCare',                'binary'),
    ('Study Identification & Design', 'Power reported or discussed',                      'Power',                      'binary'),
    ('Study Identification & Design', 'Percent retention ≥ 80%',                          'Retention (%)',              'retention80'),
    ('Study Identification & Design', 'Validated scale or measure used',                  'ValidatedScaleMeasYN',       'binary'),
    ('Study Identification & Design', 'Assessed Confounding, Moderating or Mediating',    'ConfoundYN',                 'binary'),
    ('Study Identification & Design', 'Were inclusion/exclusion criteria clearly stated?', 'InclusionExclusionCriteria', 'binary'),
    ('Study Identification & Design', 'Were funding sources and COIs disclosed?',         'Funding/COI YN',             'binary'),
    ('Population & Generalizability', 'Was baseline status reported?',                    'BaseStatusReportedYN',       'binary'),
    ('Population & Generalizability', 'Age range reported',                               'Age_Median',                 'reported'),
    ('Population & Generalizability', 'Average Percent of Participants = Female',         'PercentFem',                 'mean'),
    ('Population & Generalizability', 'Average Percent of Participants = BIPOC',          'RaceEthnicity_BIPOCPercent', 'mean'),
    ('Population & Generalizability', 'Are low SES participants included?',               'SocioeconomicStatusYN',      'binary'),
    ('Population & Generalizability', 'Country of study is the US',                       'Country_USvsNonUS',          'us'),
    ('Population & Generalizability', 'Recruitment reflective of community/clinical populations**', None,               'representative'),
    ('Intervention & Safety',         'Description of intervention content / delivery?',  'InterventionDescription',    'reported'),
    ('Intervention & Safety',         'Was therapy / counseling / integration / prep included?', 'InterventionTherapy', 'binary'),
    ('Intervention & Safety',         'Were harms or AEs assessed?',                      'HarmsAssessed',              'binary'),
    ('Equity***',                     'Was a cultural frame used',                        'CulturalFraming',            'binary'),
    ('Equity***',                     'Were Indigenous frameworks referenced?',           'IndigenousFramework',        'binary'),
    ('Equity***',                     'Did the study address reciprocity for Indigenous groups?', 'Reciprocity',        'binary'),
    ('Outcomes & Follow-Up',          'Was psychedelic intervention associated with improvement of behavior?', 'Pos Result Y/N', 'binary'),
    ('Outcomes & Follow-Up',          'Largest or Main Effect Size/Outcome reported',     'Outcome',                    'reported'),
    ('Outcomes & Follow-Up',          'P-value reported',                                 'PVal',                       'reported'),
    ('Outcomes & Follow-Up',          'Median longest follow-up in months',               'Followup in months',         'median'),
    ('Outcomes & Follow-Up',          'Were behavioral changes measured?',                'BehaviorChangeMeasured',     'binary'),
    ('Outcomes & Follow-Up',          'Were physical/biological outcomes measured?',      'PhysicalChangeMeasured',     'binary'),
    ('Outcomes & Follow-Up',          'Were mental health outcomes measured?',            'MentalHealthChangeMeasured', 'binary'),
    ('Outcomes & Follow-Up',          'Was quality of life measured?',                    'QualityOfLifeMeasured',      'binary'),
    ('Outcomes & Follow-Up',          'Were implementation outcomes measured?',           'Implementation Outcome',     'binary'),
    ('Outcomes & Follow-Up',          'Was a partial behavioral reduction reported?',     'ReductionMeasured',          'binary'),
    ('Outcomes & Follow-Up',          'Was a complete cessation reported?',               'CessationMeasured',          'binary'),
    ('Outcomes & Follow-Up',          'Was relapse tracked?',                             'RelapseMeasured',            'binary'),
    ('Outcomes & Follow-Up',          'Was mortality measured?',                          'MortalityMeasured',          'binary'),
]

FOOTNOTES = [
    '*Two studies in 1967 and 1968 matched but did not randomize.',
    '**Balancing patterns of substance use with US demographics, we filtered by '
    '≤65% male, ≥40% BIPOC, and inclusion of low-SES participants.',
    '***Reflective of USPSTF intention to expand assessment of access considerations, '
    'which can involve socio-cultural determinants, Indigenous frames, or reciprocity practices.',
]


def fmt_pct(value: float) -> str:
    """Percent, rounded half away from zero — Python's :.0f rounds half to even,
    which would print 4/32 as 12% where the manuscript convention gives 13%."""
    return f'{int(math.floor(value + 0.5))}%'


def representative(df: pd.DataFrame) -> pd.Series:
    """The manuscript footnote's composite recruitment filter."""
    female = pd.to_numeric(df['PercentFem'], errors='coerce')
    bipoc  = pd.to_numeric(df['RaceEthnicity_BIPOCPercent'], errors='coerce')
    ses    = pd.to_numeric(df['SocioeconomicStatusYN'], errors='coerce')
    return (female >= 35) & (bipoc >= 40) & (ses == 1)


def value(df: pd.DataFrame, column: str | None, kind: str) -> str:
    n = len(df)
    if n == 0:
        return '—'
    if kind in ('binary', 'reported', 'retention80'):
        return fmt_pct(da.uspstf_value(df, column, kind, n))
    if kind == 'mean':
        s = pd.to_numeric(df[column], errors='coerce').dropna()
        return '—' if s.empty else fmt_pct(s.mean())
    if kind == 'us':
        return fmt_pct(100.0 * (df[column] == 'US').sum() / n)
    if kind == 'median':
        s = pd.to_numeric(df[column], errors='coerce').dropna()
        return '—' if s.empty else f'{s.median():.0f}'
    if kind == 'representative':
        return fmt_pct(100.0 * representative(df).sum() / n)
    raise ValueError(kind)


def load_frames() -> dict[str, pd.DataFrame]:
    frames = {}
    for short, _label in CONDITIONS:
        path = _BASE / 'data' / f'{short}_Dataset_Cleaned.xlsx'
        if not path.exists():
            sys.exit(f'ERROR: {path} not found — run run_pipeline.py first.')
        frames[short] = pd.read_excel(path)
    return frames


def build_rows(frames: dict[str, pd.DataFrame]) -> tuple[list[list[str]], set[int]]:
    all_df = pd.concat(frames.values(), ignore_index=True)
    strata = [('All', all_df)] + [(label, frames[short]) for short, label in CONDITIONS]

    rows = [['Domain', 'Indicator'] + [f'{label} (n={len(df)})' for label, df in strata]]
    group_rows: set[int] = set()
    last_domain = None
    for domain, label, column, kind in INDICATORS:
        domain_cell = ''
        if domain != last_domain:
            domain_cell = domain
            group_rows.add(len(rows))
            last_domain = domain
        rows.append([domain_cell, label] + [value(df, column, kind) for _label, df in strata])
    return rows, group_rows


def main() -> None:
    frames = load_frames()
    rows, group_rows = build_rows(frames)

    doc = Document()
    da.configure_document(doc)

    title = doc.add_paragraph()
    run = title.add_run('Table 3. Methodological evaluation: '
                        'Indicators of USPSTF evidentiary standards')
    run.font.name = 'Arial'
    run.bold = True
    run.font.size = Pt(13)

    table = da.add_table(doc, n_rows=len(rows), n_cols=len(rows[0]),
                         col_widths_in=[1.5, 2.6, 0.6, 0.7, 0.65, 0.7])
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            shading = 'D9E2F3' if r == 0 else ('F2F2F2' if r in group_rows and c == 0 else None)
            da.style_cell(table.cell(r, c), text=val,
                          bold=(r == 0 or (r in group_rows and c == 0)),
                          size=9, align='left' if c < 2 else 'center', shading=shading)

    for note in FOOTNOTES:
        da.add_caption(doc, note)
    da.add_caption(doc,
                   'Missing or not-reported values count as not met (strict USPSTF gap '
                   'convention). Generated by scripts/table3_gapmap.py.', italic=True)

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')
    for row in rows:
        print(' | '.join(str(c) for c in row[1:]))


if __name__ == '__main__':
    main()
