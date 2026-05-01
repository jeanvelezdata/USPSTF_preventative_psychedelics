"""Shared constants and helpers for the USPSTF psychedelic evidence-base pipeline."""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.lines as mlines
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ── Codebook constants ────────────────────────────────────────────────────────

BINARY_COLS: list[str] = [
    'Pos Result Y/N', 'Clincial Trial YN', 'RandomizationYN', 'BlindingYN',
    'Double+Blind', 'Comparator', 'PrimaryCare', 'Power', 'ConfoundYN',
    'InclusionExclusionCriteria', 'ValidatedScaleMeasYN', 'MethodologyDescribedYN',
    'Funding/COI YN', 'SocioeconomicStatusYN', 'RecruitmentRepresentYN',
    'InterventionTherapy', 'HarmsAssessed', 'CulturalFraming',
    'IndigenousFramework', 'Reciprocity', 'BehaviorChangeMeasured',
    'PhysicalChangeMeasured', 'MentalHealthChangeMeasured', 'QualityOfLifeMeasured',
    'Implementation Outcome', 'ReductionMeasured', 'CessationMeasured',
    'RelapseMeasured', 'MortalityMeasured', 'Psilocybin', 'Ayahuasca', 'LSD',
    'Mescaline', 'Ibogaine', 'MDMA', 'Ketamine', 'Cannabis', 'OtherPsychedelic',
    'MultiplePsychedelic',
]

NUMERIC_COLS: list[str] = [
    'Retention (%)', 'PercentFem', 'RaceEthnicity_BIPOCPercent',
]

COUNTRY_MAP: dict[str, str] = {
    'USA': 'US',
    'US': 'US',
    'NonUS': 'Non-US',
    'Both': 'US and Non-US',
}

AGE_CATEGORIES: list[str] = ['<30', '30-50', '50+']
SEVERITY_CATEGORIES: list[str] = ['Mild', 'Moderate', 'Severe']

# Order used for primary_substance(); MultiplePsychedelic is handled separately.
SUBSTANCE_COLS: list[str] = [
    'Psilocybin', 'LSD', 'Mescaline', 'Ayahuasca', 'Ibogaine',
    'OtherPsychedelic', 'MDMA', 'Ketamine', 'Cannabis',
]

# Default display order in per-condition forest plots.
SUBSTANCE_ORDER: list[str] = ['Psilocybin', 'Ketamine', 'MDMA', 'LSD']

C_GRAY  = "#636363"
C_LGRAY = "#bdbdbd"

SUBSTANCE_COLORS: dict[str, str] = {
    'Psilocybin': '#1f77b4',  # blue
    'Ketamine':   '#2ca02c',  # green
    'MDMA':       '#ff7f0e',  # orange
    'LSD':        '#d62728',  # red
}

_SUBSTANCE_CLASS_MAP: dict[str, str] = {
    'Psilocybin':       'Classical psychedelic',
    'LSD':              'Classical psychedelic',
    'Mescaline':        'Classical psychedelic',
    'Ayahuasca':        'Classical psychedelic',
    'Ibogaine':         'Classical psychedelic',
    'OtherPsychedelic': 'Classical psychedelic',
    'MDMA':             'Entactogen',
    'Ketamine':         'Dissociative',
    'Cannabis':         'Cannabinoid',
}

# ── Cleaning helpers ──────────────────────────────────────────────────────────

def first_author_surname(citation: str) -> str:
    """Return the first-author surname from a citation string.

    Normalises curly apostrophes and non-breaking hyphens so config-dict keys
    can use plain ASCII.
    """
    surname = re.split(r'[,;]', str(citation).strip(), maxsplit=1)[0].strip()
    return (
        surname
        .replace('’', "'")   # right single quotation mark
        .replace('‘', "'")   # left single quotation mark
        .replace('‑', '-')   # non-breaking hyphen
    )


def recode_binary(series: pd.Series) -> pd.Series:
    """Convert binary column to {0, 1, 'Does not apply', NaN}."""
    numeric = pd.to_numeric(series, errors='coerce')
    out = pd.Series(np.nan, index=series.index, dtype=object)
    out[numeric == 0] = 0
    out[numeric == 1] = 1
    out[numeric == 3] = 'Does not apply'
    return out


def extract_effect_size_fields(
    df: pd.DataFrame, extractions: dict
) -> pd.DataFrame:
    """Add `metric` and `measure` columns from the per-condition lookup table."""
    metrics, measures = [], []
    unmatched = []
    for _, row in df.iterrows():
        key = (first_author_surname(row['Citation/Title']), int(row['Year']))
        if key in extractions:
            metric, measure, _note = extractions[key]
            metrics.append(metric if metric is not None else np.nan)
            measures.append(measure if measure is not None else np.nan)
        else:
            unmatched.append(key)
            metrics.append(np.nan)
            measures.append(np.nan)
    if unmatched:
        raise KeyError(
            f"No effect-size extraction defined for {len(unmatched)} row(s): {unmatched}. "
            "Add entries to EFFECT_SIZE_EXTRACTIONS."
        )
    df = df.copy()
    df['metric'] = metrics
    df['measure'] = measures
    cols = list(df.columns)
    cols.remove('metric')
    cols.remove('measure')
    out_idx = cols.index('Outcome')
    cols = cols[:out_idx + 1] + ['metric', 'measure'] + cols[out_idx + 1:]
    return df[cols]


def clean(df: pd.DataFrame, config: Any) -> pd.DataFrame:
    """Clean a raw extraction dataframe using per-condition config.

    Normalises column-name variants that differ across datasets, then applies
    the standard cleaning steps (whitespace, NaN replacement, numeric coercion,
    country/age/severity recoding, binary recoding, effect-size extraction).
    """
    df = df.copy()

    # Normalise column-name variants across datasets
    if 'Retention' in df.columns and 'Retention (%)' not in df.columns:
        df = df.rename(columns={'Retention': 'Retention (%)'})
    if 'StudyDesignType' in df.columns and 'Study Design Type' not in df.columns:
        df = df.rename(columns={'StudyDesignType': 'Study Design Type'})

    # 1. Strip whitespace; replace '.' with NaN (codebook convention)
    for col in df.select_dtypes(include=['object', 'string']).columns:
        df[col] = df[col].astype('string').str.strip()
        df[col] = df[col].replace({'.': pd.NA, '': pd.NA})

    # 2. Coerce designated numeric columns
    for col in NUMERIC_COLS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce')
    if 'Retention (%)' in df.columns:
        df['Retention (%)'] = df['Retention (%)'].round(1)

    # 3. Collapse country categories
    if 'Country_USvsNonUS' in df.columns:
        df['Country_USvsNonUS'] = (
            df['Country_USvsNonUS'].map(COUNTRY_MAP).fillna(df['Country_USvsNonUS'])
        )

    # 4. Age_Median: stray numeric '45' → '30-50'; ordered categorical
    if 'Age_Median' in df.columns:
        df['Age_Median'] = df['Age_Median'].astype('string').replace({'45': '30-50'})
        df['Age_Median'] = pd.Categorical(
            df['Age_Median'], categories=AGE_CATEGORIES, ordered=True
        )

    # 5. Recode binary columns: 3 → 'Does not apply'; 0/1 preserved
    for col in BINARY_COLS:
        if col in df.columns:
            df[col] = recode_binary(df[col])

    # 6. Ordered categorical for baseline severity
    if 'BaseStatus' in df.columns:
        df['BaseStatus'] = pd.Categorical(
            df['BaseStatus'], categories=SEVERITY_CATEGORIES, ordered=True
        )

    # 7. Split Outcome → metric + measure
    df = extract_effect_size_fields(df, config.EFFECT_SIZE_EXTRACTIONS)

    return df


# ── Substance classification ──────────────────────────────────────────────────

def primary_substance(row: Any) -> str:
    """Return the first positively-coded substance column (SUBSTANCE_COLS order)."""
    for col in SUBSTANCE_COLS:
        v = row.get(col) if hasattr(row, 'get') else getattr(row, col, None)
        try:
            if pd.notna(v) and float(v) == 1:
                return col
        except (TypeError, ValueError):
            continue
    return 'Unknown'


def primary_substance_class(row: Any) -> str:
    """Return the substance class per §4.3 of the project plan."""
    sub = primary_substance(row)
    if sub != 'Unknown':
        return _SUBSTANCE_CLASS_MAP.get(sub, 'Unknown')
    v = row.get('MultiplePsychedelic') if hasattr(row, 'get') else getattr(row, 'MultiplePsychedelic', None)
    try:
        if pd.notna(v) and float(v) == 1:
            return 'Multi-substance'
    except (TypeError, ValueError):
        pass
    return 'Unknown'


def classify_era(year: int) -> str:
    """Bucket a year into Pre-1980 / Gap / 2010+."""
    if year < 1980:
        return 'Pre-1980'
    if year >= 2010:
        return '2010+'
    return 'Gap'


# ── Hedges' g math ────────────────────────────────────────────────────────────

def hedges_J(dof: float) -> float:
    """Hedges' bias correction factor (Hedges & Olkin 1985 approximation)."""
    if dof < 2:
        return float('nan')
    return 1 - 3 / (4 * dof - 1)


def variance_g(g: float, n_total: float) -> float:
    """Sampling variance of Hedges' g (balanced groups assumed)."""
    n1 = n2 = n_total / 2
    return (n1 + n2) / (n1 * n2) + g**2 / (2 * (n1 + n2))


def convert_to_d(method: str, measure: Any, n_total: float, params: dict) -> float:
    """Convert a raw metric to Cohen's d."""
    m = float(measure)
    if method in ('d', 'g'):
        return m
    if method == 'F':
        return 2 * math.sqrt(m / params['df_error'])
    if method == 'eta_p':
        if m <= 0 or m >= 1:
            return float('nan')
        return 2 * math.sqrt(m / (1 - m))
    if method == 'chi2':
        phi = math.sqrt(m / n_total)
        if phi >= 1:
            phi = 0.99
        return 2 * phi / math.sqrt(1 - phi**2)
    if method in ('OR', 'RR'):
        return math.log(m) * math.sqrt(3) / math.pi
    raise ValueError(f"Unknown method: {method}")


def compute(rule: dict, measure: Any, n_total: float) -> dict | None:
    """Return d/g/var_g/se_g/ci_lo/ci_hi, or None if excluded."""
    method = rule.get('method')
    if method is None or pd.isna(measure) or pd.isna(n_total):
        return None

    direction = rule.get('direction', 1)
    d = convert_to_d(method, measure, n_total, rule) * direction

    if method == 'g':
        g = d  # already Hedges-corrected
    else:
        g = d * hedges_J(n_total - 2)

    v = variance_g(g, n_total)
    se = math.sqrt(v)
    return {
        'd': d, 'g': g, 'var_g': v, 'se_g': se,
        'ci_lo': g - 1.96 * se, 'ci_hi': g + 1.96 * se,
    }


def build_results(df: pd.DataFrame, conversion_rules: dict) -> pd.DataFrame:
    """Build the per-study Hedges' g table from a cleaned dataframe."""
    rows = []
    for _, r in df.iterrows():
        key = (first_author_surname(r['Citation/Title']), int(r['Year']))
        rule = conversion_rules.get(
            key, {'method': None, 'confidence': 'excluded', 'note': 'Rule not defined'}
        )
        n = float(r['SampleSize_NFinal'])
        result = compute(rule, r.get('measure'), n)
        rows.append({
            'study':      f"{key[0]} {key[1]}",
            'Year':       key[1],
            'N':          int(n) if not pd.isna(n) else None,
            'substance':  primary_substance(r),
            'metric':     r.get('metric'),
            'measure':    r.get('measure'),
            'method':     rule.get('method'),
            'confidence': rule.get('confidence'),
            'note':       rule.get('note', ''),
            'd':          result['d']      if result else None,
            'g':          result['g']      if result else None,
            'var_g':      result['var_g']  if result else None,
            'se_g':       result['se_g']   if result else None,
            'ci_lo':      result['ci_lo']  if result else None,
            'ci_hi':      result['ci_hi']  if result else None,
        })
    return pd.DataFrame(rows)


# ── Forest plot shared helpers ────────────────────────────────────────────────

def _build_layout(
    results: pd.DataFrame,
    substance_order: list[str] | None = None,
) -> tuple[list[dict], float, float]:
    """Filter to convertible studies, sort by substance+year, build layout list."""
    order = substance_order or SUBSTANCE_ORDER
    plot_df = results[results['g'].notna()].copy()

    known  = [s for s in order if s in plot_df['substance'].values]
    others = sorted(set(plot_df['substance'].unique()) - set(order))
    full_order = known + others

    plot_df['substance'] = pd.Categorical(
        plot_df['substance'], categories=full_order, ordered=True
    )
    plot_df = plot_df.sort_values(['substance', 'Year']).reset_index(drop=True)

    layout: list[dict] = []
    last_sub = None
    for _, r in plot_df.iterrows():
        if r['substance'] != last_sub:
            layout.append({'kind': 'header', 'text': str(r['substance'])})
            last_sub = r['substance']
        layout.append({'kind': 'study', 'data': r.to_dict()})

    if plot_df.empty:
        return layout, -1.0, 2.0
    x_min = min(float(np.min(plot_df['ci_lo'].values)), -0.5) - 0.25
    x_max = max(float(np.max(plot_df['ci_hi'].values)),  1.5) + 0.25
    return layout, x_min, x_max


def _draw_ci_row(ax: Any, r: dict, y: float) -> None:
    """Draw CI line, end caps, and square point estimate."""
    color    = SUBSTANCE_COLORS.get(r['substance'], '#444')
    is_clean = r['confidence'] == 'clean'
    alpha    = 1.0 if is_clean else 0.6
    ms       = float(np.clip(1.0 / r['var_g'] * 12, 40, 280))

    ax.plot([r['ci_lo'], r['ci_hi']], [y, y],
            color=color, alpha=alpha, lw=1.6, zorder=2, solid_capstyle='round')
    ax.plot([r['ci_lo'], r['ci_lo']], [y - 0.18, y + 0.18],
            color=color, alpha=alpha, lw=1.4)
    ax.plot([r['ci_hi'], r['ci_hi']], [y - 0.18, y + 0.18],
            color=color, alpha=alpha, lw=1.4)
    if is_clean:
        ax.scatter(r['g'], y, s=ms, color=color, marker='s',
                   edgecolor='white', lw=0.8, zorder=3)
    else:
        ax.scatter(r['g'], y, s=ms, facecolor='white',
                   edgecolor=color, marker='s', lw=1.6, zorder=3)


def _ref_lines(ax: Any, x_min: float, x_max: float) -> None:
    ax.axvline(0, color=C_GRAY, lw=0.9, linestyle='--', zorder=1)
    for b in (0.2, 0.5, 0.8):
        ax.axvline(b, color=C_LGRAY, lw=0.4, ls=':', alpha=0.6, zorder=0)


def _legend(fig: Any) -> None:
    clean_h  = mlines.Line2D([], [], color=C_GRAY, marker='s', markersize=7,
                              markerfacecolor=C_GRAY, markeredgecolor='white', lw=1.4,
                              label='Filled = direct conversion (clean)')
    approx_h = mlines.Line2D([], [], color=C_GRAY, marker='s', markersize=7,
                              markerfacecolor='white', markeredgecolor=C_GRAY, lw=1.4,
                              label='Hollow = approximate conversion')
    weight_h = mlines.Line2D([], [], lw=0, alpha=0,
                              label='Marker size ∝ inverse-variance weight  (n₁ = n₂ = N/2)')
    fig.legend(handles=[clean_h, approx_h, weight_h],
               loc='lower center', bbox_to_anchor=(0.5, 0.0),
               ncol=3, fontsize=9, framealpha=0.9, edgecolor=C_LGRAY)


# ── docx helpers ──────────────────────────────────────────────────────────────

def set_cell_shading(cell: Any, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:shd')):
        tc_pr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell: Any, color: str = 'BFBFBF', size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:tcBorders')):
        tc_pr.remove(old)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), str(size))
        b.set(qn('w:color'), color)
        borders.append(b)
    tc_pr.append(borders)


def set_cell_margins(cell: Any, top: int = 80, bottom: int = 80,
                     left: int = 120, right: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:tcMar')):
        tc_pr.remove(old)
    margins = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), str(val))
        m.set(qn('w:type'), 'dxa')
        margins.append(m)
    tc_pr.append(margins)


def style_cell(cell: Any, *, text: str | None = None, bold: bool = False,
               italic: bool = False, size: int = 10, align: str = 'left',
               shading: str | None = None, font: str = 'Arial') -> None:
    if text is not None:
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        p.alignment = {
            'left':   WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right':  WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_borders(cell)
    set_cell_margins(cell)
    if shading:
        set_cell_shading(cell, shading)


def heat_color(pct: float | None) -> str:
    """Red→amber→green gradient, lightened for readability."""
    if pct is None or (isinstance(pct, float) and math.isnan(pct)):
        return 'F2F2F2'
    pct = max(0.0, min(100.0, pct))
    if pct <= 50:
        r, g, b = 255, int(255 * (pct / 50)), 0
    else:
        r, g, b = int(255 * (1 - (pct - 50) / 50)), 200, 0
    f = 0.55
    r = int(r + (255 - r) * f)
    g = int(g + (255 - g) * f)
    b = int(b + (255 - b) * f)
    return f'{r:02X}{g:02X}{b:02X}'


def add_table(doc: Document, n_rows: int, n_cols: int,
              col_widths_in: list[float]) -> Any:
    """Create a table with explicit column widths in inches."""
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.autofit = False
    table.allow_autofit = False
    for i, w in enumerate(col_widths_in):
        for row in table.rows:
            row.cells[i].width = Inches(w)
    tbl = table._tbl
    tbl_grid = tbl.find(qn('w:tblGrid'))
    if tbl_grid is not None:
        for col, w in zip(tbl_grid.findall(qn('w:gridCol')), col_widths_in):
            col.set(qn('w:w'), str(int(w * 1440)))
            col.set(qn('w:type'), 'dxa')
    return table


def add_caption(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(9)
    run.italic = italic
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_heading(doc: Document, text: str, *, level: int = 1) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.bold = True
    run.font.size = Pt(13 if level == 1 else 11)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def configure_document(doc: Document) -> None:
    """US Letter, 1-inch margins, Arial default."""
    section = doc.sections[0]
    section.page_height = Inches(11)
    section.page_width  = Inches(8.5)
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1)
    section.right_margin  = Inches(1)
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)


def render_table(
    doc: Document,
    rows: list[list[str]],
    col_widths: list[float],
    *,
    first_row_is_header: bool = True,
    indented_rows: set[int] | None = None,
    group_rows: set[int] | None = None,
) -> None:
    """Render a generic 2-D table with header shading and section grouping."""
    indented_rows = indented_rows or set()
    group_rows    = group_rows    or set()

    table = add_table(doc, n_rows=len(rows), n_cols=len(rows[0]),
                      col_widths_in=col_widths)
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            cell = table.cell(r, c)
            shading = None
            bold    = False
            if first_row_is_header and r == 0:
                shading = 'D9E2F3'
                bold    = True
            elif r in group_rows:
                shading = 'F2F2F2'
                bold    = True
            style_cell(cell, text=str(val), bold=bold, size=10,
                       align='left' if c == 0 else 'center',
                       shading=shading)
