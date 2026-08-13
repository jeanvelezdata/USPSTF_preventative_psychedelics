"""
Manuscript Table 4 — net benefit profile among studies reporting adverse events.

Combines the two pipelines: eligible clinical trials from clinical_master.csv
and the observational corpus from observational_studies/outputs/articles.csv,
with per-event severity counts from each pipeline's ae_events file.

Inputs:
    clinical_trials/outputs/clinical_master.csv
    clinical_trials/outputs/ae_events.csv
    observational_studies/outputs/articles.csv
    observational_studies/outputs/ae_events_observational.csv

Output:
    tables/Table4_AE_Safety_Summary.docx

Usage:
    python scripts/table4_ae_safety.py
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt

_BASE    = Path(__file__).parent.parent
_SCRIPTS = Path(__file__).parent
sys.path.insert(0, str(_SCRIPTS))

import descriptive_analyses as da  # noqa: E402  (docx styling helpers)

CLINICAL_MASTER = _BASE / 'clinical_trials' / 'outputs' / 'clinical_master.csv'
CLINICAL_AE     = _BASE / 'clinical_trials' / 'outputs' / 'ae_events.csv'
OBS_ARTICLES    = _BASE / 'observational_studies' / 'outputs' / 'articles.csv'
OBS_AE          = _BASE / 'observational_studies' / 'outputs' / 'ae_events_observational.csv'
OUTPUT          = _BASE / 'tables' / 'Table4_AE_Safety_Summary.docx'

# Display order; the clinical master and the observational articles file spell
# the tobacco condition differently ('Smoking' vs both 'Smoking' and 'Tobacco').
CONDITION_ROWS = [('Tobacco', ['Smoking', 'Tobacco']),
                  ('Opioids', ['Opioid']),
                  ('Alcohol', ['Alcohol'])]

# Jones 2022 (opioid) and Jones 2022 b (tobacco) report the same 214,505-person
# sample split by outcome. Summing both would double-count those participants in
# the observational total, so the second is netted out of that row only — the
# per-condition rows each legitimately carry the sample once.
OVERLAPPING_SAMPLE = ('Jones 2022 b', 214505)

SEVERITIES = ['Mild', 'Moderate', 'Severe']


def echo(line: str) -> None:
    """Print a preview line without letting the console encoding kill the run.

    Windows consoles default to cp1252, which cannot encode the header's '≥';
    the .docx is already saved by the time we echo, so a preview character must
    never raise."""
    encoding = sys.stdout.encoding or 'utf-8'
    print(line.encode(encoding, errors='replace').decode(encoding))


def fmt_pct(value: float) -> str:
    """Percent, rounded half away from zero — Python's :.0f rounds half to even,
    which prints 5/8 as 62% where the manuscript convention gives 63%."""
    return f'{int(math.floor(value + 0.5))}%'


def n_pct(count: int, total: int) -> str:
    return f'{count} ({fmt_pct(100 * count / total)})' if total else '—'


def summarize(studies: pd.DataFrame, events: pd.DataFrame, *, label: str,
              n_col: str, positive: pd.Series, assessed: pd.Series,
              total_n: int) -> list[str]:
    """One table row. `studies` is one row per study (already de-duplicated)."""
    n_studies = len(studies)
    ev = events[events.StudyID.isin(studies.StudyID)]
    assessing = studies[assessed]
    mod_severe = set(ev.loc[ev.Severity.isin(['Moderate', 'Severe']), 'StudyID'])
    counts = ev.Severity.value_counts()
    return [
        f'{label} ({n_studies})',
        f'{total_n:,}',
        fmt_pct(100 * positive.mean()) if n_studies else '—',
        n_pct(len(assessing), n_studies),
        f'{len(set(assessing.StudyID) & mod_severe)}/{len(assessing)} '
        f'({fmt_pct(100 * len(set(assessing.StudyID) & mod_severe) / len(assessing))})'
        if len(assessing) else '—',
        *[str(int(counts.get(s, 0))) for s in SEVERITIES],
    ]


def clinical_block() -> list[list[str]]:
    master = pd.read_csv(CLINICAL_MASTER)
    events = pd.read_csv(CLINICAL_AE)
    df = master[master.EligibleYN == 'Y']

    def row(label: str, sub: pd.DataFrame) -> list[str]:
        return summarize(
            sub, events, label=label, n_col='SampleSize_NFinal',
            positive=pd.to_numeric(sub['Pos Result Y/N'], errors='coerce') == 1,
            assessed=pd.to_numeric(sub['HarmsAssessed'], errors='coerce') == 1,
            total_n=int(pd.to_numeric(sub['SampleSize_NFinal'], errors='coerce').sum()),
        )

    rows = [row('Clinical Trial', df)]
    for label, names in CONDITION_ROWS:
        rows.append(row('    ' + label, df[df.Condition.isin(names)]))
    return rows


def observational_block() -> list[list[str]]:
    articles = pd.read_csv(OBS_ARTICLES)
    events = pd.read_csv(OBS_AE)
    # The corpus is curated to be all-observational, but the extraction can mark
    # a study IsObservational=N (merge_validate flags it as a non-blocking
    # warning). Such a study must not reach the observational counts.
    articles = articles[articles.IsObservational == 'Y']

    def row(label: str, sub: pd.DataFrame, *, net_overlap: bool = False) -> list[str]:
        unique = sub.drop_duplicates('StudyID')
        total = int(pd.to_numeric(unique.N, errors='coerce').sum())
        if net_overlap:
            name, n = OVERLAPPING_SAMPLE
            if name in set(unique.StudyID):
                total -= n
        return summarize(
            unique, events, label=label, n_col='N',
            positive=unique.PosResult.astype('string').str.upper().eq('Y'),
            assessed=unique.SafetyReported.eq('Y'),
            total_n=total,
        )

    rows = [row('Observational', articles, net_overlap=True)]
    for label, names in CONDITION_ROWS:
        rows.append(row('    ' + label, articles[articles.Condition.isin(names)]))
    return rows


def main() -> None:
    for path in (CLINICAL_MASTER, CLINICAL_AE, OBS_ARTICLES, OBS_AE):
        if not path.exists():
            sys.exit(f'ERROR: {path} not found.')

    body = clinical_block() + observational_block()
    header = [
        ['Study Type (N)', 'N total', 'Studies with positive SU outcome',
         'Studies report assessing AEs',
         'Studies with ≥1 moderate/severe AE (of those assessing)',
         'AE events by severity', '', ''],
        ['', '', '', '', '', 'Mild', 'Mod.', 'Severe'],
    ]
    rows = header + body
    group_rows = {2, 6}  # the two totals rows

    doc = Document()
    da.configure_document(doc)

    title = doc.add_paragraph()
    run = title.add_run('Table 4. Net benefit profile among studies reporting '
                        'on adverse events (AEs)')
    run.font.name = 'Arial'
    run.bold = True
    run.font.size = Pt(13)

    table = da.add_table(doc, n_rows=len(rows), n_cols=8,
                         col_widths_in=[1.35, 0.85, 0.95, 0.9, 1.1, 0.5, 0.5, 0.55])
    table.cell(0, 5).merge(table.cell(0, 7))
    for col in range(5):
        table.cell(0, col).merge(table.cell(1, col))
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            if r == 1 and c < 5:
                continue          # merged upward
            if r == 0 and 5 < c:
                continue          # merged leftward
            is_header = r < 2
            da.style_cell(table.cell(r, c), text=str(val),
                          bold=is_header or r in group_rows, size=9,
                          align='left' if c == 0 else 'center',
                          shading='D9E2F3' if is_header else
                                  ('F2F2F2' if r in group_rows else None))

    da.add_caption(doc, 'AE = adverse event. Severity counts are discrete AE events, '
                        'not participants.')
    da.add_caption(doc, f'Observational N total nets out the {OVERLAPPING_SAMPLE[1]:,}-person '
                        f'sample shared by Jones 2022 and {OVERLAPPING_SAMPLE[0]}, which the '
                        f'per-condition rows each count once. '
                        f'Generated by scripts/table4_ae_safety.py.', italic=True)

    doc.save(OUTPUT)
    print(f'Wrote {OUTPUT}')
    for row in rows:
        echo(' | '.join(str(c) for c in row))


if __name__ == '__main__':
    main()
